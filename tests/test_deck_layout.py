"""
Deck layout QA — geometry, not content.

A pptx that builds is not a pptx that reads. Nothing in pptxgenjs stops a text
box running off the slide or a table column landing on top of the next one, and
the failure is invisible until it is on a screen in front of an investor. This
checks the geometry directly.

It exists because a phase description on the roadmap slide ran into the tranche
label beside it and the file built without complaint.

Run:  python3 tests/test_deck_layout.py [path/to/deck.pptx]
"""

from __future__ import annotations

import glob
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from pptx import Presentation  # noqa: E402
from pptx.util import Emu  # noqa: E402

BLEED_TOL = Emu(int(0.02 * 914400))     # 0.02in of rounding slack at the edges

# Text boxes are laid out taller than their text so a line can wrap without
# reflowing the slide, so adjacent boxes routinely share a hairline of padding.
# Flagging those buries the real collisions in noise. A collision counts when
# the shared rectangle is a meaningful share of the smaller box AND is deep
# enough to actually put glyphs on glyphs.
MIN_OVERLAP_FRACTION = 0.20
MIN_OVERLAP_DEPTH = Emu(int(0.12 * 914400))


def _rect(sh):
    if sh.left is None or sh.top is None or sh.width is None or sh.height is None:
        return None
    return (sh.left, sh.top, sh.left + sh.width, sh.top + sh.height)


def _has_text(sh) -> bool:
    return sh.has_text_frame and bool(sh.text_frame.text.strip())


def _overlap(a, b) -> tuple[int, int]:
    return (min(a[2], b[2]) - max(a[0], b[0]), min(a[3], b[3]) - max(a[1], b[1]))


def check(path: Path) -> list[str]:
    prs = Presentation(str(path))
    W, H = prs.slide_width, prs.slide_height
    problems: list[str] = []

    for i, slide in enumerate(prs.slides, start=1):
        boxes = []
        for sh in slide.shapes:
            r = _rect(sh)
            if r is None:
                continue
            if r[0] < -BLEED_TOL or r[1] < -BLEED_TOL or r[2] > W + BLEED_TOL \
                    or r[3] > H + BLEED_TOL:
                problems.append(
                    f"slide {i}: '{sh.shape_type}' runs off the slide "
                    f"({r[0]/914400:.2f},{r[1]/914400:.2f})–"
                    f"({r[2]/914400:.2f},{r[3]/914400:.2f}) on "
                    f"{W/914400:.2f}x{H/914400:.2f}")
            if _has_text(sh):
                boxes.append((sh, r))

        # Two text boxes stacked on each other is the failure mode that matters.
        # Backing panels are excluded above by the has-text filter.
        for a in range(len(boxes)):
            for b in range(a + 1, len(boxes)):
                sh_a, ra = boxes[a]
                sh_b, rb = boxes[b]
                dx, dy = _overlap(ra, rb)
                if dx <= 0 or dy <= 0:
                    continue
                smaller = min((ra[2] - ra[0]) * (ra[3] - ra[1]),
                              (rb[2] - rb[0]) * (rb[3] - rb[1]))
                frac = (dx * dy) / smaller if smaller else 0.0
                if frac > MIN_OVERLAP_FRACTION and dy > MIN_OVERLAP_DEPTH:
                    ta = sh_a.text_frame.text.strip().replace("\n", " ")[:38]
                    tb = sh_b.text_frame.text.strip().replace("\n", " ")[:38]
                    problems.append(
                        f"slide {i}: text overlaps by "
                        f"{dx/914400:.2f}x{dy/914400:.2f}in ({frac:.0%} of the "
                        f"smaller box) — '{ta}' / '{tb}'")

    return problems


def _self_test() -> list[str]:
    """
    Tuning a checker until it reports nothing is the same as deleting it. Build
    a deck with a deliberate collision and a deliberate bleed, and confirm both
    are still caught at the current tolerances.
    """
    import tempfile

    from pptx import Presentation as P
    from pptx.util import Inches

    prs = P()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    sl.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = "left"
    sl.shapes.add_textbox(Inches(2), Inches(1), Inches(4), Inches(1)).text_frame.text = "right"
    sl.shapes.add_textbox(Inches(12), Inches(7), Inches(4), Inches(1)).text_frame.text = "bleed"
    with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as fh:
        prs.save(fh.name)
        found = check(Path(fh.name))
    missing = []
    if not any("overlaps" in p for p in found):
        missing.append("self-test: a 3in text collision was NOT caught")
    if not any("runs off the slide" in p for p in found):
        missing.append("self-test: a shape past the right edge was NOT caught")
    return missing


def main() -> int:
    if len(sys.argv) > 1:
        paths = [Path(sys.argv[1])]
    else:
        paths = [Path(p) for p in sorted(glob.glob(str(ROOT / "dist" / "*.pptx")))]
    if not paths:
        print("no deck found — build one first")
        return 1

    failed = 0
    for problem in _self_test():
        print(f"  {problem}")
        failed += 1

    for path in paths:
        prs = Presentation(str(path))
        problems = check(path)
        print(f"\n{path.name}: {len(prs.slides)} slides")
        for p in problems:
            print(f"  {p}")
        failed += len(problems)
        if not problems:
            print("  no geometry problems")

    print(f"\n{'PASS — deck geometry clean' if not failed else f'{failed} LAYOUT PROBLEM(S)'}")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
